from docx import Document
from openpyxl import load_workbook

from valcopilot import ibbi
from valcopilot.build import build
from valcopilot.report import indian


def test_indian_grouping():
    assert indian(1234567.891) == "12,34,567.89"
    assert indian(999.5) == "999.50"
    assert indian(-100000, 0) == "-1,00,000"
    assert indian(57.1449) == "57.14"


def test_report_has_every_ibbi_section_in_order(sample_xlsx, tmp_path):
    result = build(sample_xlsx, tmp_path / "r.docx")
    assert result.errors == []
    doc = Document(result.report)
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 2" and ") " in p.text[:3]]
    assert headings == [f"{s.letter}) {s.title}" for s in ibbi.SECTIONS]


def test_vrin_in_the_footer_and_values_in_the_summary(sample_xlsx, tmp_path):
    result = build(sample_xlsx, tmp_path / "r.docx")
    doc = Document(result.report)
    assert doc.sections[0].footer.paragraphs[0].text.startswith("VRIN: SAMPLE-VRIN-0000")
    summary = {row.cells[0].text: row.cells[1].text for row in doc.tables[0].rows}
    assert list(summary)[1:] == [f.label for f in ibbi.EXEC_SUMMARY]
    from valcopilot.report import cr
    assert summary["Fair Value"] == f"INR {cr(result.conclusion.fair_value)} crore"
    assert summary["Class of Asset"] == "Securities or Financial Assets"
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "ILLUSTRATIVE SAMPLE" in text and "WORKING DRAFT" not in text


def test_open_issues_are_flagged_on_the_title_page(sample_xlsx, tmp_path):
    wb = load_workbook(sample_xlsx)
    ws = wb["Engagement"]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "conflict")
    ws.cell(row, 3).value = None
    wb.save(sample_xlsx)
    result = build(sample_xlsx, tmp_path / "r.docx")
    text = "\n".join(p.text for p in Document(result.report).paragraphs)
    assert "WORKING DRAFT - 1 open checklist issue(s)" in text
    assert "[MISSING" in text


def test_strict_refuses_to_build_with_errors(sample_xlsx, tmp_path):
    wb = load_workbook(sample_xlsx)
    ws = wb["Engagement"]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "vrin")
    ws.cell(row, 3).value = None
    wb.save(sample_xlsx)
    result = build(sample_xlsx, tmp_path / "r.docx", strict=True)
    assert result.report is None and result.errors


def test_valuation_error_stops_the_build(sample_xlsx, tmp_path):
    wb = load_workbook(sample_xlsx)
    wb["Approaches"]["D2"] = 0.5   # weights now 90%
    wb.save(sample_xlsx)
    result = build(sample_xlsx, tmp_path / "r.docx")
    assert result.report is None
    assert any("not 100%" in str(i) for i in result.errors)


def test_draft_facts_label_each_figure(sample_inputs):
    # The live run used 58.40 (100% equity) as the stake value, the subsidiary's history as the
    # corporate debtor's, and said "four identified, one excluded" when five peers were considered.
    from valcopilot.engine import conclude
    from valcopilot.report import facts_for
    c = conclude(sample_inputs)
    vr = facts_for("value_rationale", sample_inputs, c)
    assert "100% of the equity" in vr["concluded_equity"] and "58.40" in vr["concluded_equity"]
    assert "74%" in vr["stake_value"] and "43.21" in vr["stake_value"] and "58.40" in vr["stake_value"]
    assert "100% of the equity" in vr["dcf"] and "45.38" in vr["dcf"]
    m = facts_for("approaches_text", sample_inputs, c)
    assert m["comps"].startswith("5 comparable companies considered; 1 excluded (Illustrative Peer 5")
    assert "4 used" in m["comps"]
    bg = facts_for("background", sample_inputs, c)
    assert "Sample Packaging" in bg["valued_entity_history"]
    assert "not of the corporate debtor" in bg["valued_entity_history"]


def test_draft_facts_count_the_approaches_applied(sample_inputs):
    # The second live run said "three valuation approaches have been applied"; only two were.
    from valcopilot.engine import conclude
    from valcopilot.report import facts_for
    f = facts_for("value_rationale", sample_inputs, conclude(sample_inputs))
    assert f["approaches"].startswith("3 approaches considered; 2 applied (Income, Market); 1 not applied (Cost).")
