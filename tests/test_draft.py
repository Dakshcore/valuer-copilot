"""AI drafting, tested with a fake client. No API calls are made here."""
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import load_workbook

from valcopilot.build import build
from valcopilot.draft import draft_section, make_client, unknown_numbers


class FakeClient:
    def __init__(self, text):
        self.text, self.calls = text, []
        self.messages = self

    def create(self, **kw):
        self.calls.append(kw)
        return SimpleNamespace(content=[SimpleNamespace(type="text", text=self.text)],
                               usage=SimpleNamespace(input_tokens=500, output_tokens=120))


def test_number_guard():
    facts = "fair value 57.14, stake 74%, 74,00,000 shares, FY26"
    assert unknown_numbers("Fair value is INR 57.14 crore for 7400000 shares (74%) in FY26.", facts) == []
    assert unknown_numbers("Fair value is 57.140 crore.", facts) == []           # trailing zero
    assert unknown_numbers("Fair value is 58.2 crore, up 12%.", facts) == ["12", "58.2"]


def test_draft_uses_only_the_given_facts():
    client = FakeClient("The corporate debtor holds a stake valued at INR 57.14 crore.")
    d = draft_section("summary", {"total_fair_value": "57.14"}, client=client, model="m")
    assert d.ok and d.input_tokens == 500
    call = client.calls[0]
    assert call["model"] == "m" and "Use only the facts provided" in call["system"]
    assert "57.14" in call["messages"][0]["content"]


def test_invented_number_is_flagged():
    d = draft_section("summary", {"total_fair_value": "57.14"}, client=FakeClient("Value is INR 60.00 crore."))
    assert not d.ok and d.unknown_numbers == ["60"]


def test_missing_key_gives_a_clear_error(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    with pytest.raises(RuntimeError, match="ANTHROPIC_API_KEY"):
        make_client()


def blank(xlsx, *keys):
    wb = load_workbook(xlsx)
    ws = wb["Engagement"]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value in keys:
            ws.cell(r, 3).value = None
    wb.save(xlsx)


def test_build_drafts_only_blank_sections_and_marks_them(sample_xlsx, tmp_path):
    blank(sample_xlsx, "summary", "background")
    client = FakeClient("Drafted text for the section.")
    result = build(sample_xlsx, tmp_path / "r.docx", draft=True, client=client)
    assert set(result.drafts) == {"summary", "background"}
    assert result.errors == []
    text = "\n".join(p.text for p in Document(result.report).paragraphs)
    assert text.count("AI-drafted text. The valuer must review") == 2


def test_build_flags_a_draft_with_an_invented_number(sample_xlsx, tmp_path):
    blank(sample_xlsx, "summary")
    result = build(sample_xlsx, tmp_path / "r.docx", draft=True, client=FakeClient("Total value 99.99 crore."))
    assert any("99.99" in str(i) for i in result.errors)
    text = "\n".join(p.text for p in Document(result.report).paragraphs)
    assert "Number check failed: 99.99" in text
