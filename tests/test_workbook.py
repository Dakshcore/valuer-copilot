import pytest
from openpyxl import load_workbook

from valcopilot.workbook import InputError, _pct, read_workbook, write_workbook


def test_sample_round_trip(sample_inputs, sample_data):
    assert sample_inputs.fields["vrin"] == sample_data["fields"]["vrin"]
    assert sample_inputs.assumptions.dlom == pytest.approx(0.15)
    assert sample_inputs.assumptions.mid_year is True
    assert sample_inputs.assumptions.shares == 20_000_000
    assert [y.projected for y in sample_inputs.years] == [False] * 3 + [True] * 5
    assert sample_inputs.years[3].pat is None
    assert len(sample_inputs.comps) == 5 and len(sample_inputs.receivables) == 6
    assert sample_inputs.assets[0].stake == pytest.approx(0.74)
    assert sample_inputs.receivables[4].disputed is True


def test_template_is_empty_but_readable(tmp_path):
    inp = read_workbook(write_workbook(tmp_path / "t.xlsx"))
    assert inp.assumptions is None and inp.assets == [] and inp.get("vrin") == ""


def test_percent_parsing():
    assert _pct(0.25, "x") == 0.25      # cell formatted as %
    assert _pct(25, "x") == 0.25        # typed as a plain number
    assert _pct("25%", "x") == 0.25
    assert _pct(1, "x") == 1.0          # 100%
    assert _pct(None, "x", blank=0.0) == 0.0


def test_bad_row_names_the_sheet_and_row(sample_xlsx):
    wb = load_workbook(sample_xlsx)
    wb["Financials"]["B3"] = "Forecast"
    wb.save(sample_xlsx)
    with pytest.raises(InputError, match="Financials row 3"):
        read_workbook(sample_xlsx)


def test_non_number_is_reported(sample_xlsx):
    wb = load_workbook(sample_xlsx)
    wb["Comps"]["B2"] = "five hundred"
    wb.save(sample_xlsx)
    with pytest.raises(InputError, match="Comps row 2 Market Cap"):
        read_workbook(sample_xlsx)


def test_dates_typed_as_excel_dates_are_read_as_dd_mm_yyyy(sample_xlsx):
    from datetime import datetime
    wb = load_workbook(sample_xlsx)
    ws = wb["Engagement"]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "icd")
    ws.cell(row, 3).value = datetime(2026, 4, 1)
    wb.save(sample_xlsx)
    assert read_workbook(sample_xlsx).fields["icd"] == "01/04/2026"
