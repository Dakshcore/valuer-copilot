"""The Excel input workbook: write a blank template or a filled sample, and read one back."""
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from . import ibbi
from .model import Approach, Asset, Assumptions, Comp, Inputs, Receivable, Year
from .sample import SAMPLE

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")

ENGAGEMENT_FIELDS = [f for f in ibbi.EXEC_SUMMARY if f.key not in ("class_of_asset", "fair_value", "liquidation_value")] \
    + ibbi.BODY_FIELDS + ibbi.ANNEXURES

# key, label, guidance, kind ("pct", "num", "bool", "text")
ASSUMPTIONS = [
    ("risk_free", "Risk-free rate", "e.g. 10-year G-sec yield, cite the source in section o", "pct"),
    ("beta", "Beta", "Relevered beta of the subject", "num"),
    ("erp", "Equity risk premium", "", "pct"),
    ("specific_premium", "Size / company-specific risk premium", "", "pct"),
    ("cost_of_debt", "Pre-tax cost of debt", "", "pct"),
    ("tax_rate", "Tax rate", "", "pct"),
    ("debt_weight", "Debt / (Debt + Equity)", "Target capital structure", "pct"),
    ("terminal_growth", "Terminal growth rate", "", "pct"),
    ("mid_year", "Mid-year discounting (Yes/No)", "", "bool"),
    ("debt", "Debt, INR crore", "As on the valuation date", "num"),
    ("cash", "Cash and equivalents, INR crore", "As on the valuation date", "num"),
    ("shares", "Number of equity shares", "Units, not crore", "num"),
    ("dlom", "Discount for lack of marketability", "Applied to the concluded equity value; 0 if none", "pct"),
    ("comps_multiple", "Comps multiple used", "EV/EBITDA, EV/Revenue or P/E", "text"),
    ("comps_stat", "Comps statistic", "median or mean", "text"),
    ("cost_value", "Adjusted net asset value, INR crore", "Only for the cost approach", "num"),
]

TABLES = {
    "Financials": ["Year", "Type (Actual/Projected)", "Revenue", "EBITDA", "D&A", "Capex",
                   "Increase in NWC", "PAT"],
    "Comps": ["Company", "Market Cap", "Net Debt", "Revenue", "EBITDA", "PAT"],
    "Approaches": ["Approach", "Applicable (Yes/No)", "Method", "Weight %", "Reason for applying or not applying"],
    "Assets": ["Asset", "Category", "Premise of value", "Book Value", "Basis (Valuation/Receivables/Input)",
               "Stake %", "Fair Value (Input basis only)", "Liquidation value as % of fair value", "Remarks"],
    "Receivables": ["Debtor", "Nature", "Related party (Yes/No)", "Secured (Yes/No)", "Disputed (Yes/No)",
                    "Age bucket", "Amount", "FV recovery %", "LV recovery %", "Remarks"],
}
PCT_COLUMNS = {"Approaches": {4}, "Assets": {6, 8}, "Receivables": {8, 9}}

README = [
    "IBC Valuation Report Builder - Securities or Financial Assets",
    f"Report format per {ibbi.CIRCULAR}.",
    "All amounts are INR crore unless a label says otherwise. Yellow cells are inputs.",
    "Sheets: Engagement (summary fields and report text), Financials, Assumptions, Comps, Approaches, Assets, Receivables.",
    "Text marked 'N/A' must give a reason, e.g. 'N/A: no expert was involved' (Note 2 of the format).",
    "Run:  python -m valcopilot check <this file>   then   python -m valcopilot build <this file> -o report.docx",
]


def _header(ws: Worksheet, cols: list[str]) -> None:
    ws.append(cols)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def write_workbook(path: Path, sample: bool = False) -> Path:
    data = SAMPLE if sample else {}
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    for line in README + (["", "THIS FILE CONTAINS FICTIONAL SAMPLE DATA FOR DEMONSTRATION ONLY."] if sample else []):
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Engagement")
    _header(ws, ["Key", "Field", "Value", "Guidance"])
    for f in ENGAGEMENT_FIELDS:
        ws.append([f.key, f.label, data.get("fields", {}).get(f.key, ""), f.hint + ("" if f.required else " (optional)")])
        ws.cell(ws.max_row, 3).fill = INPUT_FILL
    _style_kv(ws)

    ws = wb.create_sheet("Assumptions")
    _header(ws, ["Key", "Item", "Value", "Guidance"])
    values = data.get("assumptions", {})
    for key, label, hint, kind in ASSUMPTIONS:
        v = values.get(key, "")
        if kind == "bool" and v != "":
            v = "Yes" if v else "No"
        ws.append([key, label, v, hint])
        cell = ws.cell(ws.max_row, 3)
        cell.fill = INPUT_FILL
        if kind == "pct":
            cell.number_format = "0.00%"
    _style_kv(ws)

    for name, cols in TABLES.items():
        ws = wb.create_sheet(name)
        _header(ws, cols)
        for row in data.get(name.lower(), []):
            ws.append(list(row))
        for r in range(2, max(ws.max_row, 30) + 1):
            for c in range(1, len(cols) + 1):
                cell = ws.cell(r, c)
                cell.fill = INPUT_FILL
                if c in PCT_COLUMNS.get(name, ()):
                    cell.number_format = "0.00%"
        for c in range(1, len(cols) + 1):
            ws.column_dimensions[ws.cell(1, c).column_letter].width = 18 if c > 1 else 34
    path = Path(path)
    wb.save(path)
    return path


def _style_kv(ws: Worksheet) -> None:
    ws.column_dimensions["A"].hidden = True
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 70
    ws.column_dimensions["D"].width = 55
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


# ---------- reading ----------

class InputError(ValueError):
    pass


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _num(v, where: str, blank=None):
    if v is None or (isinstance(v, str) and not v.strip()):
        if blank is not None:
            return blank
        raise InputError(f"{where}: a number is required.")
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").replace("%", "").strip())
    except ValueError:
        raise InputError(f"{where}: {v!r} is not a number.") from None


def _pct(v, where: str, blank=None):
    """Percent cells hold 0.25 when formatted as %. A plain number above 1 (e.g. 25) is read as 25%."""
    if isinstance(v, str) and v.strip().endswith("%"):
        return _num(v, where) / 100
    x = _num(v, where, blank)
    return x / 100 if x is not None and abs(x) > 1 else x


def _bool(v, where: str) -> bool:
    t = _text(v).lower()
    if t in ("yes", "y", "true", "1"):
        return True
    if t in ("no", "n", "false", "0", ""):
        return False
    raise InputError(f"{where}: expected Yes or No, got {v!r}.")


def _rows(ws: Worksheet):
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row and any(c not in (None, "") for c in row):
            yield i, row


def read_workbook(path: Path) -> Inputs:
    wb = load_workbook(path, data_only=True)
    missing = [s for s in ("Engagement", "Assumptions", *TABLES) if s not in wb.sheetnames]
    if missing:
        raise InputError(f"Missing sheet(s): {', '.join(missing)}. Start from 'python -m valcopilot template'.")
    inp = Inputs()
    for _, row in _rows(wb["Engagement"]):
        if row[0]:
            inp.fields[str(row[0]).strip()] = _text(row[2] if len(row) > 2 else None)

    kv = {str(r[0]).strip(): (r[2] if len(r) > 2 else None) for _, r in _rows(wb["Assumptions"]) if r[0]}
    if any(_text(v) for v in kv.values()):
        args = {}
        for key, label, _, kind in ASSUMPTIONS:
            where = f"Assumptions!{label}"
            v = kv.get(key)
            if kind == "pct":
                args[key] = _pct(v, where, blank=0.0 if key in ("dlom", "specific_premium") else None)
            elif kind == "bool":
                args[key] = _bool(v, where)
            elif kind == "num":
                args[key] = _num(v, where, blank=None if key != "cost_value" else float("nan"))
            else:
                args[key] = _text(v) or {"comps_multiple": "EV/EBITDA", "comps_stat": "median"}[key]
        if args["cost_value"] != args["cost_value"]:  # NaN means blank
            args["cost_value"] = None
        inp.assumptions = Assumptions(**args)

    for i, r in _rows(wb["Financials"]):
        w = f"Financials row {i}"
        kind = _text(r[1]).lower()
        if kind not in ("actual", "projected", "a", "p"):
            raise InputError(f"{w}: Type must be Actual or Projected.")
        inp.years.append(Year(_text(r[0]), kind.startswith("p"), _num(r[2], w + " Revenue"), _num(r[3], w + " EBITDA"),
                              _num(r[4], w + " D&A", 0.0), _num(r[5], w + " Capex", 0.0),
                              _num(r[6], w + " NWC", 0.0), _num(r[7], w + " PAT", None) if r[7] not in (None, "") else None))
    for i, r in _rows(wb["Comps"]):
        w = f"Comps row {i}"
        inp.comps.append(Comp(_text(r[0]), _num(r[1], w + " Market Cap"), _num(r[2], w + " Net Debt", 0.0),
                              _num(r[3], w + " Revenue"), _num(r[4], w + " EBITDA"), _num(r[5], w + " PAT")))
    for i, r in _rows(wb["Approaches"]):
        w = f"Approaches row {i}"
        applicable = _bool(r[1], w)
        inp.approaches.append(Approach(_text(r[0]).title(), applicable, _text(r[2]),
                                       _pct(r[3], w + " Weight", 0.0) if applicable else 0.0, _text(r[4])))
    for i, r in _rows(wb["Assets"]):
        w = f"Assets row {i}"
        basis = _text(r[4]).title()
        inp.assets.append(Asset(_text(r[0]), _text(r[1]), _text(r[2]), _num(r[3], w + " Book Value", 0.0), basis,
                                _pct(r[5], w + " Stake", 1.0),
                                _num(r[6], w + " Fair Value", None) if basis == "Input" else None,
                                _pct(r[7], w + " Liquidation %", 1.0), _text(r[8] if len(r) > 8 else None)))
    for i, r in _rows(wb["Receivables"]):
        w = f"Receivables row {i}"
        inp.receivables.append(Receivable(_text(r[0]), _text(r[1]), _bool(r[2], w), _bool(r[3], w), _bool(r[4], w),
                                          _text(r[5]), _num(r[6], w + " Amount"), _pct(r[7], w + " FV recovery"),
                                          _pct(r[8], w + " LV recovery"), _text(r[9] if len(r) > 9 else None)))
    return inp
