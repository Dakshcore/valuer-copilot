"""Write the valuation report as a Word document in the IBBI SFA format."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from . import ibbi
from .draft import Draft
from .engine import Conclusion, sensitivity
from .model import Inputs

MISSING = "[MISSING: to be completed by the valuer]"
REVIEW_COLOUR = RGBColor(0xC0, 0x00, 0x00)


# ---------- number formats (also used to build the facts given to the AI drafter) ----------

def indian(x: float, decimals: int = 2) -> str:
    """Indian digit grouping: 12,34,567.89"""
    neg, x = x < 0, abs(x)
    whole, _, frac = f"{x:.{decimals}f}".partition(".")
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        groups = []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        whole = ",".join(([head] if head else []) + groups + [tail])
    return ("-" if neg else "") + whole + (f".{frac}" if decimals else "")


def cr(x: float | None) -> str:
    return "-" if x is None else indian(x)


def pct(x: float | None, decimals: int = 2) -> str:
    return "-" if x is None else f"{x * 100:.{decimals}f}%"


def mult(x: float | None) -> str:
    return "n.m." if x is None else f"{x:.2f}x"


# ---------- docx helpers ----------

def _table(doc, header: list[str], rows: list[list[str]], widths: list[float] | None = None, bold_last=False,
           size: float | None = None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        _shade(cell, "D9E2F3")
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, (cell, text) in enumerate(zip(cells, row)):
            cell.text = str(text)
            if j > 0 and _looks_numeric(str(text)):
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if bold_last and i == len(rows) - 1:
                for r in cell.paragraphs[0].runs:
                    r.bold = True
    if widths:
        t.autofit = False
        for col, w in zip(t.columns, widths):
            col.width = Cm(w)
        for row in t.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Cm(w)
    if size:
        for row in t.rows:
            for cell in row.cells:
                for r in cell.paragraphs[0].runs:
                    r.font.size = Pt(size)
    doc.add_paragraph()
    return t


def _looks_numeric(s: str) -> bool:
    return bool(s) and s.replace(",", "").replace(".", "").replace("%", "").replace("x", "").replace("-", "").isdigit()


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _page_field(paragraph) -> None:
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def _text(doc, text: str) -> None:
    for para in (text or MISSING).split("\n"):
        if para.strip():
            doc.add_paragraph(para.strip())


def _bullets(doc, text: str) -> None:
    lines = [ln.strip(" -•\t") for ln in (text or "").split("\n") if ln.strip()]
    if not lines:
        doc.add_paragraph(MISSING)
    for ln in lines:
        doc.add_paragraph(ln, style="List Bullet")


def _review_note(doc, d: Draft | None) -> None:
    if d is None:
        return
    p = doc.add_paragraph()
    run = p.add_run("AI-drafted text. The valuer must review and edit it before signing.")
    run.italic, run.font.size, run.font.color.rgb = True, Pt(8.5), REVIEW_COLOUR
    if not d.ok:
        run = p.add_run(f" Number check failed: {', '.join(d.unknown_numbers)} do not appear in the inputs.")
        run.bold, run.font.size, run.font.color.rgb = True, Pt(8.5), REVIEW_COLOUR


# ---------- the report ----------

def write_report(path: Path, inp: Inputs, c: Conclusion, drafts: dict[str, Draft] | None = None,
                 open_issues: int = 0, sample: bool = False) -> Path:
    drafts = drafts or {}
    f = inp.get
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.2))
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)

    footer = sec.footer.paragraphs[0]
    # The Footer style carries a centre tab; clear it so one tab reaches the right edge.
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CLEAR)
    footer.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT)
    footer.add_run(f"VRIN: {f('vrin') or MISSING}\tPage ").font.size = Pt(8.5)
    _page_field(footer)

    # Title page
    for _ in range(5):
        doc.add_paragraph()
    for text, size, bold in (("VALUATION REPORT", 24, True), (ibbi.CLASS_OF_ASSET, 16, False),
                             (f("corporate_debtor").split(",")[0] or MISSING, 14, True),
                             ("Under the Insolvency and Bankruptcy Code, 2016", 12, False),
                             (f"Valuation date: {f('valuation_date') or MISSING}", 11, False),
                             (f"Prepared in the format specified by {ibbi.CIRCULAR}", 9, False)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size, run.bold = Pt(size), bold
    for banner in (["ILLUSTRATIVE SAMPLE - FICTIONAL DATA - NOT A REAL VALUATION"] if sample else []) + \
                  ([f"WORKING DRAFT - {open_issues} open checklist issue(s)"] if open_issues else []):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(banner)
        run.bold, run.font.color.rgb = True, REVIEW_COLOUR
    doc.add_page_break()

    # Executive summary
    doc.add_heading("Executive Summary of the Valuation Report", level=1)
    computed = {"class_of_asset": ibbi.CLASS_OF_ASSET, "fair_value": f"INR {cr(c.fair_value)} crore",
                "liquidation_value": f"INR {cr(c.liquidation_value)} crore"}
    _table(doc, ["Item", "Details"], [[fl.label, computed.get(fl.key) or f(fl.key) or MISSING]
                                      for fl in ibbi.EXEC_SUMMARY], [6, 10.5])

    s = {x.letter: x for x in ibbi.SECTIONS}
    h = lambda letter: doc.add_heading(f"{letter}) {s[letter].title}", level=2)  # noqa: E731

    h("a"); _text(doc, s["a"].guidance); _text(doc, f("scope_detail"))
    h("b")
    _text(doc, f"{f('valuer_name') or MISSING} - {ibbi.CLASS_OF_ASSET}\n"
               f"Registration Number: {f('registration_no') or MISSING}\n"
               f"VRIN: {f('vrin') or MISSING}. The VRIN is shown in the footer on the left side of each page.")
    h("c"); _text(doc, f("experts"))
    h("d"); _text(doc, f("conflict"))
    h("e")
    _text(doc, f"Client/Appointing Authority: {f('client') or MISSING}\n"
               f"Other intended users: {f('other_users') or 'None'}")
    h("f"); _text(doc, s["f"].guidance)

    h("g")
    _text(doc, f"Subject asset(s): {f('subject_asset') or MISSING}")
    _table(doc, ["Asset", "Category", "Book value (INR crore)", "Basis of valuation", "Remarks"],
           [[a.name, a.category, cr(a.book_value), a.basis, a.remarks] for a in inp.assets], [5, 3, 2.5, 2.5, 3.5])

    h("h"); _text(doc, f("background")); _review_note(doc, drafts.get("background"))

    h("i")
    currency = f("currency") or "INR"
    _text(doc, f"Valuation currency: {currency}. Amounts are in INR crore unless stated otherwise.")
    if currency.upper() != "INR":
        _text(doc, f"Exchange rate to INR: {f('fx_rate') or MISSING} as on {f('fx_date') or MISSING}, "
                   f"source: {f('fx_source') or MISSING}.")

    h("j"); _bullets(doc, f("sources_info"))

    h("k")
    doc.add_paragraph("Basis Adopted", style="Heading 3")
    _table(doc, ["Valuation Purpose", "Basis"],
           [["Determination of fair value under the IBC", "Fair value"],
            ["Determination of liquidation value under the IBC", "Liquidation value"]], [10, 6.5])
    _text(doc, s["k"].guidance)
    doc.add_paragraph("Premise Adopted", style="Heading 3")
    _table(doc, ["Asset Type", "Premise Adopted", "Remarks/Justification"],
           [[a.name, a.premise or MISSING, a.remarks or "-"] for a in inp.assets], [6, 4, 6.5])

    h("l"); _text(doc, f("standards"))

    h("m")
    _text(doc, f("approaches_text")); _review_note(doc, drafts.get("approaches_text"))
    _table(doc, ["Approach", "Applicable", "Method adopted", "Weight", "Reason"],
           [[ap.name, "Yes" if ap.applicable else "No", ap.method or "-", pct(ap.weight, 0) if ap.applicable else "-",
             ap.reason or MISSING] for ap in inp.approaches], [2.2, 2, 4, 1.8, 6.5])
    doc.add_paragraph("Meetings with the Committee of Creditors", style="Heading 3")
    _text(doc, f("coc_meetings"))

    h("n")
    rows = []
    if c.dlom:
        rows.append(["Discount for lack of marketability", "Concluded equity value of the valued entity", pct(c.dlom)])
    for a in inp.assets:
        if a.basis != "Receivables" and a.lv_recovery < 1:
            rows.append(["Liquidation discount", a.name, pct(1 - a.lv_recovery)])
    if c.receivables:
        rows.append(["Receivable recovery haircuts", "Trade receivables", "See Annexure B"])
    if rows:
        _table(doc, ["Adjustment", "Applied to", "Rate"], rows, [5, 8.5, 3])
    else:
        _text(doc, "No discounts or premiums were applied.")
    _text(doc, f("discounts_rationale"))

    h("o"); _bullets(doc, f("data_inputs"))
    if c.comps:
        _text(doc, "Comparable company data and the multiples derived from it are set out in Annexure A.")
    h("p"); _bullets(doc, f("inspections"))
    h("q"); _text(doc, f("sustainability")); _review_note(doc, drafts.get("sustainability"))

    h("r")
    _bullets(doc, f("assumptions"))
    if c.dcf:
        a = inp.assumptions
        doc.add_paragraph("Model conventions", style="Heading 3")
        _bullets(doc, "The valuation date is taken as the start of the first projected year.\n"
                      + ("Cash flows are discounted from the middle of each year.\n" if a.mid_year
                         else "Cash flows are discounted from the end of each year.\n")
                      + "The terminal value uses the Gordon growth method on the final year's free cash flow "
                        "and is discounted from the end of the final projected year.")

    h("s")
    _table(doc, ["Asset", "Book value", "Fair value", "Liquidation value"],
           [[v.name, cr(v.book_value), cr(v.fair_value), cr(v.liquidation_value)] for v in c.assets]
           + [["Total (INR crore)", cr(c.book_value), cr(c.fair_value), cr(c.liquidation_value)]],
           [7.5, 3, 3, 3], bold_last=True)
    if c.approach_values:
        doc.add_paragraph("Value of the valued entity's equity by approach", style="Heading 3")
        rows = [[name, method, cr(v), pct(w, 0), cr(v * w)] for name, method, v, w in c.approach_values]
        rows += [["Weighted equity value", "", "", "", cr(c.weighted_equity)],
                 ["Less: discount for lack of marketability", "", "", pct(c.dlom), cr(c.weighted_equity - c.concluded_equity)],
                 ["Concluded equity value (INR crore)", "", "", "", cr(c.concluded_equity)]]
        _table(doc, ["Approach", "Method", "Equity value", "Weight", "Weighted"], rows, [4.5, 4.5, 2.5, 2, 3],
               bold_last=True)
    if f("replacement_cost"):
        _text(doc, f"Replacement cost: INR {f('replacement_cost')} crore.")
    _text(doc, f("value_rationale")); _review_note(doc, drafts.get("value_rationale"))

    h("t"); _text(doc, f("caveats"))
    h("u")
    _text(doc, f("summary")); _review_note(doc, drafts.get("summary"))
    _text(doc, f"Fair value: INR {cr(c.fair_value)} crore. Liquidation value: INR {cr(c.liquidation_value)} crore. "
               f"Book value: INR {cr(c.book_value)} crore.")

    h("v")
    rows = [[an.label, f(an.key) or (MISSING if an.required else "-")] for an in ibbi.ANNEXURES]
    if c.comps:
        rows.append(["Annexure A: Comparable company analysis", "Attached"])
    rows.append(["Annexure B: Computation sheets", "Attached"])
    _table(doc, ["Annexure", "Status"], rows, [10, 6.5])

    _annexures(doc, inp, c)
    path = Path(path)
    doc.save(path)
    return path


def _annexures(doc, inp: Inputs, c: Conclusion) -> None:
    a = inp.assumptions
    if c.comps:
        doc.add_page_break()
        doc.add_heading("Annexure A: Comparable Company Analysis", level=1)
        _text(doc, "INR crore. Multiples are not meaningful (n.m.) where the denominator is zero or negative.")
        _table(doc, ["Company", "EV/Revenue", "EV/EBITDA", "P/E"],
               [[name, mult(v["EV/Revenue"]), mult(v["EV/EBITDA"]), mult(v["P/E"])] for name, v in c.comps.table],
               [7, 3, 3, 3])
        _table(doc, ["Statistic", "EV/Revenue", "EV/EBITDA", "P/E"],
               [[stat.title(), *(mult(c.comps.summary.get(m, {}).get(stat)) for m in ("EV/Revenue", "EV/EBITDA", "P/E"))]
                for stat in ("min", "median", "mean", "max")], [7, 3, 3, 3])
        rows = [[f"Selected {c.comps.multiple} ({c.comps.stat})", mult(c.comps.selected)],
                ["Subject metric (latest actual year)", cr(c.comps.subject_metric)]]
        if c.comps.enterprise_value is not None:
            rows += [["Enterprise value", cr(c.comps.enterprise_value)], ["Less: debt", cr(a.debt)],
                     ["Add: cash", cr(a.cash)]]
        rows += [["Equity value", cr(c.comps.equity_value)], ["Value per share (INR)", indian(c.comps.per_share)]]
        if c.comps.excluded:
            rows.append(["Excluded (multiple not meaningful)", ", ".join(c.comps.excluded)])
        _table(doc, ["Item", "Value"], rows, [10, 6.5])

    doc.add_page_break()
    doc.add_heading("Annexure B: Computation Sheets", level=1)
    if c.dcf:
        d = c.dcf
        doc.add_heading("B1. Cost of capital", level=2)
        _table(doc, ["Input", "Value"],
               [["Risk-free rate", pct(a.risk_free)], ["Beta", f"{a.beta:.2f}"], ["Equity risk premium", pct(a.erp)],
                ["Size / company-specific premium", pct(a.specific_premium)], ["Cost of equity", pct(d.ke)],
                ["Pre-tax cost of debt", pct(a.cost_of_debt)], ["Tax rate", pct(a.tax_rate)],
                ["Debt / (Debt + Equity)", pct(a.debt_weight)], ["WACC", pct(d.wacc)],
                ["Terminal growth", pct(a.terminal_growth)]], [10, 6.5])
        doc.add_heading("B2. Discounted cash flow (INR crore)", level=2)
        proj = [y for y in inp.years if y.projected]
        _table(doc, ["Year", "Revenue", "EBITDA", "D&A", "Capex", "Incr. NWC", "FCFF", "Period", "Factor", "PV"],
               [[y.label, cr(y.revenue), cr(y.ebitda), cr(y.da), cr(y.capex), cr(y.nwc_change), cr(r.fcff),
                 f"{r.period:.1f}", f"{r.factor:.4f}", cr(r.pv)] for y, r in zip(proj, d.rows)], size=9)
        _table(doc, ["Item", "Value"],
               [["Sum of PV of FCFF", cr(sum(r.pv for r in d.rows))], ["Terminal value", cr(d.terminal_value)],
                ["PV of terminal value", cr(d.pv_terminal)], ["Enterprise value", cr(d.enterprise_value)],
                ["Less: debt", cr(a.debt)], ["Add: cash", cr(a.cash)], ["Equity value", cr(d.equity_value)],
                ["Value per share (INR)", indian(d.per_share)]], [10, 6.5])
        doc.add_heading("B3. Sensitivity of equity value (INR crore): WACC vs terminal growth", level=2)
        rates, growths, grid = sensitivity(inp.years, a)
        _table(doc, ["WACC \\ g", *(pct(g) for g in growths)],
               [[pct(r), *(cr(v) for v in row)] for r, row in zip(rates, grid)])
    if c.receivables:
        rec = c.receivables
        doc.add_heading("B4. Receivables (IBBI Part I(c))", level=2)
        _table(doc, ["Debtor", "Nature", "Related", "Secured", "Disputed", "Age", "Amount", "FV %", "LV %", "Remarks"],
               [[r.debtor, r.nature, "Yes" if r.related_party else "No", "Yes" if r.secured else "No",
                 "Yes" if r.disputed else "No", r.age, cr(r.amount), pct(r.fv_recovery, 0), pct(r.lv_recovery, 0),
                 r.remarks or "-"] for r in inp.receivables],
               [1.9, 2.0, 1.5, 1.5, 1.6, 1.7, 1.6, 1.2, 1.2, 2.4], size=8.5)
        _table(doc, ["Age bucket", "Gross", "Fair value", "Liquidation value"],
               [[age, cr(g), cr(fv), cr(lv)] for age, (g, fv, lv) in rec.by_age.items()]
               + [["Total", cr(rec.gross), cr(rec.fair_value), cr(rec.liquidation_value)]], bold_last=True)


# ---------- facts handed to the AI drafter ----------

def facts_for(key: str, inp: Inputs, c: Conclusion) -> dict[str, str]:
    f = inp.get
    facts = {
        "corporate_debtor": f("corporate_debtor"),
        "subject_asset": f("subject_asset"),
        "purpose": f("purpose"),
        "valuation_date": f("valuation_date"),
        "currency_and_units": "INR crore",
        "corporate_debtor_assets": "; ".join(f"{v.name} ({v.category}): book value {cr(v.book_value)}, fair value {cr(v.fair_value)}, "
                            f"liquidation value {cr(v.liquidation_value)}" for v in c.assets),
        "total_fair_value": cr(c.fair_value),
        "total_liquidation_value": cr(c.liquidation_value),
        "total_book_value": cr(c.book_value),
    }
    asset_notes = "; ".join(f"{a.name}: {a.remarks}" for a in inp.assets if a.remarks)
    # Label every figure with what it measures; an unlabelled figure gets attached to the wrong item.
    held = next((a for a in inp.assets if a.basis == "Valuation"), None)
    entity = f"the company whose shares are valued (asset: {held.name})" if held else "the valued company"
    whole = f"100% of the equity of {entity}"
    if key in ("background", "sustainability"):
        facts["asset_remarks"] = asset_notes
        facts["background_notes_from_valuer"] = f("background")
        actual = [y for y in inp.years if not y.projected]
        if actual:
            facts["valued_entity_history"] = (f"Historical financials of {entity}, not of the corporate debtor: "
                                              + "; ".join(f"{y.label}: revenue {cr(y.revenue)}, EBITDA {cr(y.ebitda)}"
                                                          for y in actual))
        if c.receivables:
            r = c.receivables
            facts["receivables"] = (f"gross {cr(r.gross)} across {len(inp.receivables)} debtors; by age: "
                                    + "; ".join(f"{age} {cr(g)}" for age, (g, _, _) in r.by_age.items())
                                    + f"; disputed debtors: {sum(x.disputed for x in inp.receivables)}; "
                                      f"related-party debtors: {sum(x.related_party for x in inp.receivables)}")
    if key in ("approaches_text", "value_rationale", "summary"):
        used = [ap.name for ap in inp.approaches if ap.applicable]
        unused = [ap.name for ap in inp.approaches if not ap.applicable]
        facts["approaches"] = (f"{len(inp.approaches)} approaches considered; {len(used)} applied ({', '.join(used)})"
                               + (f"; {len(unused)} not applied ({', '.join(unused)}). " if unused else ". ")
                               ) + "; ".join(
            f"{ap.name}: {'applied' if ap.applicable else 'not applied'}"
            + (f", method {ap.method}, weight {pct(ap.weight, 0)}" if ap.applicable else "") + f". Reason: {ap.reason}"
            for ap in inp.approaches)
    if key in ("approaches_text", "value_rationale", "summary") and c.dcf:
        facts["dcf"] = (f"WACC {pct(c.dcf.wacc)}, terminal growth {pct(inp.assumptions.terminal_growth)}. "
                        f"DCF value of {whole}, before discount for lack of marketability: {cr(c.dcf.equity_value)}")
    if key in ("approaches_text", "value_rationale", "summary") and c.comps:
        k = c.comps
        metric = {"P/E": "PAT"}.get(k.multiple, k.multiple.split("/")[-1])
        excluded = (f"{len(k.excluded)} excluded ({', '.join(k.excluded)}: {metric} zero or negative, so the "
                    f"multiple is not meaningful)" if k.excluded else "none excluded")
        facts["comps"] = (f"{len(k.table)} comparable companies considered; {excluded}; {len(k.used)} used. "
                          f"{k.stat} {k.multiple} of the companies used: {mult(k.selected)}. "
                          f"Comparable companies value of {whole}, before discount for lack of marketability: "
                          f"{cr(k.equity_value)}")
    if key in ("value_rationale", "summary") and c.concluded_equity is not None:
        facts["concluded_equity"] = (f"Value of {whole}: weighted across approaches {cr(c.weighted_equity)}; "
                                     f"after discount for lack of marketability of {pct(c.dlom)}: "
                                     f"{cr(c.concluded_equity)}")
        stake = next((v for v in c.assets if held and v.name == held.name), None)
        if stake:
            facts["stake_value"] = (f"Corporate debtor's {pct(held.stake, 0)} stake: fair value {cr(stake.fair_value)} "
                                    f"({pct(held.stake, 0)} of {cr(c.concluded_equity)}), liquidation value "
                                    f"{cr(stake.liquidation_value)}")
        facts["discounts_rationale"] = f("discounts_rationale")
    return {k: v for k, v in facts.items() if v}
